#!/usr/bin/env python3
"""Build the Selected-50 research-only Market Intelligence workbook.

No prediction model is run here. The three supported raw Oracle probabilities are
copied from the runtime CSV without recalculation. Every archive join is exact and
fail-closed on missing or duplicate composite identities.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from bisect import bisect_left
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = [
    "Fixture Date", "League", "Fixture", "BTTS Yes", "Home 1X", "Away X2",
    "Over 1.5", "Over 2.5", "Under 4.5", "Home over 1.5 goals",
    "Away over 1.5 goals", "Home 1X + Over 1.5", "Away X2 + Over 1.5",
    "Home 1X + Under 4.5", "Away X2 + Under 4.5", "M-1UP Home",
    "M-1UP Away", "Draw Probability (derived)", "Goal Line Call (O/U 2.5)",
    "Most Likely Total-Goals Band (derived)", "Recent Form (5)",
    "Recent Form Points (Last 5)", "League Position",
    "Avg. Goals Scored (Last 5)", "Avg. Goals Conceded (Last 5)",
    "Avg. Points (Last 5)", "Attacking Index", "Defensive Index",
    "TOTAL Index", "H2H Last 5 Summary", "H2H Avg Total Goals",
    "Historical Accuracy Status", "Home Evidence Risk", "Away Evidence Risk",
    "Home Risk Rationale", "Away Risk Rationale", "A+ Setup Status",
    "Gem's Independent Match Analysis",
]
SUPPORTED = {
    "BTTS Yes": "BTTS_probability_pct",
    "Home 1X": "DC_HOME_probability_pct",
    "Away X2": "DC_AWAY_probability_pct",
}
UNAVAILABLE = "NOT_AVAILABLE_FROM_ORACLE_OUTPUT"
SETTLED = {"complete", "completed", "finished", "ft", "full time", "full_time",
           "played", "after extra time", "aet", "penalties"}


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fields):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def parse_dt(value):
    try:
        d = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if d.tzinfo is None:
            d = d.replace(tzinfo=timezone.utc)
        return d.astimezone(timezone.utc)
    except (TypeError, ValueError):
        return None


def valid_score(row):
    try:
        h, a = float(row.get("home_goals", "")), float(row.get("away_goals", ""))
        return (math.isfinite(h) and math.isfinite(a) and h >= 0 and a >= 0
                and h.is_integer() and a.is_integer())
    except (TypeError, ValueError, OverflowError):
        return False


def archive_key(row):
    return (
        str(row.get("provider_season_id") or "").strip(),
        str(row.get("fixture_id") or "").strip(),
        str(row.get("home_team_id") or "").strip(),
        str(row.get("away_team_id") or "").strip(),
    )


def runtime_key(row):
    return (
        str(row.get("provider_season_id") or "").strip(),
        str(row.get("fixture_id") or "").strip(),
        str(row.get("home_team_provider_id") or "").strip(),
        str(row.get("away_team_provider_id") or "").strip(),
    )


def build_indexes(matches):
    """Create strictly chronological archive indexes from exact provider identities."""
    by_team = defaultdict(list)
    by_season = defaultdict(list)
    by_pair = defaultdict(list)
    seen = set()
    duplicates = 0
    for row in matches:
        dt = parse_dt(row.get("kickoff_utc"))
        if not dt or not row.get("home_team_id") or not row.get("away_team_id"):
            continue
        season_id = str(row.get("provider_season_id") or "").strip()
        league = season_id
        season = season_id
        h, a = str(row["home_team_id"]), str(row["away_team_id"])
        if not league or not season or h == a:
            continue
        identity = (str(row.get("provider_season_id") or ""), str(row.get("fixture_id") or ""), h, a)
        if all(identity) and identity in seen:
            duplicates += 1
        elif all(identity):
            seen.add(identity)
        item = dict(row)
        item["_dt"] = dt
        item["_league"] = league
        item["_season"] = season
        item["_home"] = h
        item["_away"] = a
        by_team[(league, h)].append(item)
        by_team[(league, a)].append(item)
        by_season[(league, season)].append(item)
        by_pair[(league, tuple(sorted((h, a))))].append(item)
    for index in (by_team, by_season, by_pair):
        for key in index:
            index[key].sort(key=lambda x: (x["_dt"], str(x.get("fixture_id") or "")))
    return by_team, by_season, by_pair, duplicates


def before(rows, cutoff):
    times = [r["_dt"] for r in rows]
    return rows[:bisect_left(times, cutoff)]


def team_form(rows, team_id):
    usable = [r for r in rows if str(r.get("status_raw") or "").lower() in SETTLED
              and valid_score(r)]
    sample = usable[-5:]
    results, gf, ga, pts = [], [], [], []
    for r in sample:
        home = str(r.get("home_team_id")) == team_id
        own = int(float(r["home_goals"] if home else r["away_goals"]))
        opp = int(float(r["away_goals"] if home else r["home_goals"]))
        gf.append(own); ga.append(opp)
        points = 3 if own > opp else 1 if own == opp else 0
        pts.append(points)
        results.append("W" if points == 3 else "D" if points == 1 else "L")
    n = len(sample)
    return {
        "n": n, "form": "".join(results),
        "gf": sum(gf) / n if n else None,
        "ga": sum(ga) / n if n else None,
        "ppg": sum(pts) / n if n else None,
        "points": sum(pts) if n else None,
        "sample": sample,
    }


def league_table(rows, cutoff):
    stats = {}
    for r in rows:
        if r["_dt"] >= cutoff:
            break
        if str(r.get("status_raw") or "").lower() not in SETTLED or not valid_score(r):
            continue
        h, a = str(r["_home"]), str(r["_away"])
        hg, ag = int(float(r["home_goals"])), int(float(r["away_goals"]))
        for tid, gf, ga in ((h, hg, ag), (a, ag, hg)):
            team_name = str((r.get("home_team") if tid == h else r.get("away_team")) or "").strip()
            s = stats.setdefault(tid, {"played": 0, "points": 0, "gf": 0, "ga": 0, "team_name": team_name})
            if team_name and not s.get("team_name"):
                s["team_name"] = team_name
            s["played"] += 1; s["gf"] += gf; s["ga"] += ga
            s["points"] += 3 if gf > ga else 1 if gf == ga else 0
    ordered = sorted(stats.items(), key=lambda kv: (
        -kv[1]["points"], -(kv[1]["gf"] - kv[1]["ga"]), -kv[1]["gf"],
        str(kv[1].get("team_name") or "").casefold(), kv[0]
    ))
    return {tid: {"position": i + 1, **s} for i, (tid, s) in enumerate(ordered)}


def pct(value):
    if value in (None, ""):
        return ""
    return round(float(value), 2)


def index_value(team_gf, team_ga, league_goals_per_team):
    if team_gf is None or team_ga is None or league_goals_per_team is None:
        return None, None
    attack = min(100.0, max(0.0, 50.0 * team_gf / league_goals_per_team))
    defence = None if team_ga == 0 else min(100.0, max(0.0, 50.0 * league_goals_per_team / team_ga))
    return attack, defence


def exact_join(runtime_rows, archive_rows):
    index = defaultdict(list)
    for row in archive_rows:
        key = archive_key(row)
        if all(key):
            index[key].append(row)
    duplicates = {k: len(v) for k, v in index.items() if len(v) != 1}
    joined, missing = {}, []
    for row in runtime_rows:
        key = runtime_key(row)
        if not all(key):
            missing.append((row, "MISSING_COMPOSITE_IDENTITY"))
        elif len(index.get(key, [])) != 1:
            missing.append((row, "NO_EXACT_ARCHIVE_MATCH" if not index.get(key) else "DUPLICATE_ARCHIVE_MATCH"))
        else:
            joined[key] = index[key][0]
    return joined, missing, duplicates


def build_report(runtime_rows, archive_rows):
    joined, missing, duplicate_archive_keys = exact_join(runtime_rows, archive_rows)
    by_team, by_season, by_pair, history_duplicates = build_indexes(archive_rows)
    output, trace = [], []
    leagues = set()
    seen_runtime = set()
    for runtime in runtime_rows:
        key = runtime_key(runtime)
        selected = str(runtime.get("selected_id") or "").strip()
        league_name = str(runtime.get("league") or runtime.get("provider_competition_name_source") or "")
        leagues.add(selected)
        if key in seen_runtime:
            raise ValueError(f"Duplicate runtime exact key: {key}")
        seen_runtime.add(key)
        target = joined.get(key)
        cutoff = parse_dt(runtime.get("date_utc"))
        if target is None or cutoff is None:
            reason = next((why for row, why in missing if row is runtime), "MISSING_ARCHIVE_JOIN_OR_KICKOFF")
            outrow = base_row(runtime, league_name, reason)
            output.append(outrow)
            trace.append({"fixture_id": runtime.get("fixture_id", ""), "selected_id": selected,
                          "join_status": reason, "home_form_n": 0, "away_form_n": 0,
                          "home_table_status": "UNAVAILABLE", "away_table_status": "UNAVAILABLE",
                          "history_cutoff_utc": runtime.get("date_utc", "")})
            continue
        league = key[0]
        season = key[0]
        home, away = key[2], key[3]
        home_rows = before(by_team[(league, home)], cutoff)
        away_rows = before(by_team[(league, away)], cutoff)
        hf, af = team_form(home_rows, home), team_form(away_rows, away)
        table = league_table(by_season[(league, season)], cutoff)
        hpos, apos = table.get(home, {}).get("position"), table.get(away, {}).get("position")
        league_completed = [r for r in before(by_season[(league, season)], cutoff)
                            if str(r.get("status_raw") or "").lower() in SETTLED and valid_score(r)]
        total_goals = sum(int(float(r["home_goals"])) + int(float(r["away_goals"])) for r in league_completed)
        league_goals_per_team = (total_goals / (2 * len(league_completed))) if league_completed else None
        h_attack, h_defence = index_value(hf["gf"] if hf["n"] >= 3 else None,
                                           hf["ga"] if hf["n"] >= 3 else None, league_goals_per_team)
        a_attack, a_defence = index_value(af["gf"] if af["n"] >= 3 else None,
                                           af["ga"] if af["n"] >= 3 else None, league_goals_per_team)
        pair_rows = before(by_pair[(league, tuple(sorted((home, away))))], cutoff)
        h2h = [r for r in pair_rows if str(r.get("status_raw") or "").lower() in SETTLED and valid_score(r)][-5:]
        h2h_results = []
        for r in h2h:
            home_was_home = str(r.get("home_team_id")) == home
            own = int(float(r["home_goals"] if home_was_home else r["away_goals"]))
            opp = int(float(r["away_goals"] if home_was_home else r["home_goals"]))
            h2h_results.append("W" if own > opp else "D" if own == opp else "L")
        h2h_avg = (sum(int(float(r["home_goals"])) + int(float(r["away_goals"])) for r in h2h) / len(h2h)) if h2h else None
        h_risk, h_reasons = evidence_risk(hf["n"], hpos is not None, len(h2h), all(key))
        a_risk, a_reasons = evidence_risk(af["n"], apos is not None, len(h2h), all(key))
        h_attack_text, h_def_text = fmt(h_attack), fmt(h_defence)
        a_attack_text, a_def_text = fmt(a_attack), fmt(a_defence)
        h_total = h_attack + h_defence if h_attack is not None and h_defence is not None else None
        a_total = a_attack + a_defence if a_attack is not None and a_defence is not None else None
        def pairfmt(x, y):
            return f"Home {fmt(x)} | Away {fmt(y)}" if x is not None or y is not None else "INSUFFICIENT_DATA"
        form_text = f"Home {hf['form'] or 'INSUFFICIENT'} | Away {af['form'] or 'INSUFFICIENT'}"
        points_text = (f"Home {hf['points']}/{hf['n'] * 3} (n={hf['n']})" if hf["points"] is not None else "Home N/A") + " | " + (f"Away {af['points']}/{af['n'] * 3} (n={af['n']})" if af["points"] is not None else "Away N/A")
        position_text = f"Home {hpos if hpos else 'N/A'} | Away {apos if apos else 'N/A'} (PROVISIONAL)"
        avg_gf = pairfmt(hf["gf"] if hf["n"] >= 3 else None, af["gf"] if af["n"] >= 3 else None)
        avg_ga = pairfmt(hf["ga"] if hf["n"] >= 3 else None, af["ga"] if af["n"] >= 3 else None)
        avg_pts = pairfmt(hf["ppg"] if hf["n"] >= 3 else None, af["ppg"] if af["n"] >= 3 else None)
        attack_text = pairfmt(h_attack, a_attack)
        defence_text = pairfmt(h_defence, a_defence)
        total_text = pairfmt(h_total, a_total)
        h2h_text = "".join(h2h_results) if h2h_results else "NO_VERIFIED_PRIOR_H2H"
        analysis = independent_analysis(hf, af, h_attack, a_attack, h_defence, a_defence)
        row = base_row(runtime, league_name, "EXACT_ARCHIVE_JOIN")
        row.update({
            "Recent Form (5)": form_text,
            "Recent Form Points (Last 5)": points_text,
            "League Position": position_text,
            "Avg. Goals Scored (Last 5)": avg_gf,
            "Avg. Goals Conceded (Last 5)": avg_ga,
            "Avg. Points (Last 5)": avg_pts,
            "Attacking Index": attack_text,
            "Defensive Index": defence_text,
            "TOTAL Index": total_text,
            "H2H Last 5 Summary": h2h_text,
            "H2H Avg Total Goals": fmt(h2h_avg),
            "Historical Accuracy Status": "NOT_VALIDATED_NO_EXACT_PREMATCH_LEDGER_AND_OFFICIAL_SETTLEMENT_JOIN",
            "Home Evidence Risk": h_risk,
            "Away Evidence Risk": a_risk,
            "Home Risk Rationale": h_reasons,
            "Away Risk Rationale": a_reasons,
            "A+ Setup Status": "NOT_GRADED_CALIBRATION_SETTLEMENT_AND_ODDS_GATES_NOT_PASSED",
            "Gem's Independent Match Analysis": analysis,
        })
        output.append(row)
        trace.append({
            "fixture_id": runtime.get("fixture_id", ""), "selected_id": selected,
            "join_status": "PASS_EXACT_SEASON_FIXTURE_ORDERED_TEAM_IDS",
            "home_form_n": hf["n"], "away_form_n": af["n"], "h2h_n": len(h2h),
            "home_league_position": hpos or "", "away_league_position": apos or "",
            "league_table_status": "PROVISIONAL_CONVENTIONAL_TABLE" if hpos and apos else "TABLE_POSITION_INSUFFICIENT_OR_UNAVAILABLE",
            "history_cutoff_utc": cutoff.isoformat(),
            "history_cutoff_rule": "STRICTLY_BEFORE_FIXTURE_KICKOFF",
            "home_risk_pct": h_risk, "away_risk_pct": a_risk,
            "archive_provider_season_id": key[0], "archive_fixture_id": key[1],
            "archive_home_team_id": key[2], "archive_away_team_id": key[3],
        })
    if len(output) != len(runtime_rows):
        raise AssertionError("Output row count does not reconcile to runtime fixture rows")
    return output, trace, {
        "runtime_rows": len(runtime_rows), "report_rows": len(output),
        "exact_archive_joins": sum(r.get("join_status") == "PASS_EXACT_SEASON_FIXTURE_ORDERED_TEAM_IDS" for r in trace),
        "unmatched_or_quarantined_rows": sum(r.get("join_status") != "PASS_EXACT_SEASON_FIXTURE_ORDERED_TEAM_IDS" for r in trace),
        "selected_league_groups": len(leagues), "selected_ids": sorted(leagues),
        "duplicate_archive_composite_keys": len(duplicate_archive_keys),
        "duplicate_archive_history_rows": history_duplicates,
        "duplicate_runtime_keys": len(runtime_rows) - len(seen_runtime),
    }


def fmt(value):
    return "N/A" if value is None else f"{value:.2f}"


def evidence_risk(form_n, has_position, h2h_n, identity_ok):
    # Evidence coverage score only; not a probability of losing or match outcome.
    components = {
        "form_gap": max(0.0, (5 - min(form_n, 5)) / 5) * 40,
        "standings_gap": 0.0 if has_position else 30.0,
        "h2h_gap": max(0.0, (5 - min(h2h_n, 5)) / 5) * 20,
        "identity_gap": 0.0 if identity_ok else 10.0,
    }
    risk = round(sum(components.values()), 1)
    rationale = ("Evidence coverage gap: " + ", ".join(
        f"{k}={v:.0f}/" + str({"form_gap":40,"standings_gap":30,"h2h_gap":20,"identity_gap":10}[k])
        for k, v in components.items() if v > 0
    )) if risk else "Required evidence coverage present; standings remain provisional."
    return risk, rationale


def independent_analysis(hf, af, ha, aa, hd, ad):
    if hf["n"] < 3 or af["n"] < 3:
        return "Insufficient recent five-match evidence for a balanced comparison; no independent match conclusion assigned."
    parts = []
    if hf["ppg"] > af["ppg"]:
        parts.append("Home side has the stronger recent points rate")
    elif hf["ppg"] < af["ppg"]:
        parts.append("Away side has the stronger recent points rate")
    else:
        parts.append("Recent points rates are level")
    if ha is not None and aa is not None:
        parts.append("home attacking index is higher" if ha > aa else "away attacking index is higher" if aa > ha else "attacking indices are level")
    if hd is not None and ad is not None:
        parts.append("home defensive index is higher" if hd > ad else "away defensive index is higher" if ad > hd else "defensive indices are level")
    return "; ".join(parts) + ". Evidence summary only; not a calibrated prediction or betting recommendation."


def base_row(runtime, league_name, join_status):
    row = {h: "" for h in HEADERS}
    row["Fixture Date"] = runtime.get("date_utc", "")
    row["League"] = league_name
    row["Fixture"] = f"{runtime.get('home_team','')} vs {runtime.get('away_team','')}"
    for header, source in SUPPORTED.items():
        value = runtime.get(source, "")
        row[header] = value if value not in (None, "") else "NOT_AVAILABLE_FROM_ORACLE_OUTPUT"
    for header in HEADERS[6:20]:
        row[header] = UNAVAILABLE
    row["Historical Accuracy Status"] = "NOT_VALIDATED"
    row["A+ Setup Status"] = "NOT_GRADED"
    if join_status != "EXACT_ARCHIVE_JOIN":
        row["Recent Form (5)"] = join_status
        row["Recent Form Points (Last 5)"] = join_status
        row["League Position"] = join_status
        row["Avg. Goals Scored (Last 5)"] = join_status
        row["Avg. Goals Conceded (Last 5)"] = join_status
        row["Avg. Points (Last 5)"] = join_status
        row["Attacking Index"] = join_status
        row["Defensive Index"] = join_status
        row["TOTAL Index"] = join_status
        row["H2H Last 5 Summary"] = join_status
        row["H2H Avg Total Goals"] = join_status
        row["Home Evidence Risk"] = "UNQUANTIFIED_IDENTITY_OR_ARCHIVE_JOIN_FAILURE"
        row["Away Evidence Risk"] = "UNQUANTIFIED_IDENTITY_OR_ARCHIVE_JOIN_FAILURE"
        row["Home Risk Rationale"] = join_status
        row["Away Risk Rationale"] = join_status
        row["A+ Setup Status"] = "NOT_GRADED"
        row["Gem's Independent Match Analysis"] = "No analysis: exact archive identity join failed."
    return row


def write_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Market Intelligence"
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in HEADERS])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="333366")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    widths = [14, 20, 34] + [18] * 17 + [32] * 18
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--runtime-report", required=True)
    ap.add_argument("--archive-audit", required=True, help="Directory containing match_field_audit.csv")
    ap.add_argument("--schedule-coverage", required=True, help="Runtime schedule coverage CSV proving the full governed league universe")
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--expected-league-groups", type=int, default=50)
    args = ap.parse_args()
    out_dir = Path(args.out_dir); out_dir.mkdir(parents=True, exist_ok=True)
    runtime_rows = read_csv(args.runtime_report)
    archive_rows = read_csv(Path(args.archive_audit) / "match_field_audit.csv")
    schedule_rows = read_csv(args.schedule_coverage)
    governed_ids = {str(r.get("selected_id") or "").strip() for r in schedule_rows if str(r.get("selected_id") or "").strip()}
    report_ids = {str(r.get("selected_id") or "").strip() for r in runtime_rows if str(r.get("selected_id") or "").strip()}
    if len(schedule_rows) != args.expected_league_groups or len(governed_ids) != args.expected_league_groups:
        raise SystemExit(f"FAIL: governed schedule coverage must contain exactly {args.expected_league_groups} unique selected IDs")
    if not report_ids.issubset(governed_ids):
        raise SystemExit("FAIL: runtime report contains selected IDs outside the governed 50-league schedule")
    report, trace, summary = build_report(runtime_rows, archive_rows)
    fixture_league_groups = summary["selected_league_groups"]
    summary["fixture_league_groups"] = fixture_league_groups
    summary["selected_league_groups"] = len(governed_ids)
    summary["selected_ids_without_fixture_rows_in_runtime_horizon"] = sorted(governed_ids - report_ids)
    summary["selected_ids_with_fixture_rows_in_runtime_horizon"] = sorted(report_ids)
    summary.update({
        "schema_status": "PASS_38_COLUMNS" if len(HEADERS) == 38 else "FAIL_SCHEMA",
        "headers": HEADERS,
        "production_eligible": False, "betting_eligible": False,
        "production_model_logic_changed": False, "production_probabilities_changed": False,
        "probability_policy": "BTTS, DC_HOME and DC_AWAY copied unchanged from runtime CSV; all other probability/market fields explicitly unavailable.",
        "combo_policy": "Combo probabilities remain masked; no combo output is synthesized.",
        "historical_accuracy_status": "NOT_VALIDATED",
        "standings_policy": "Provisional conventional 3-1-0 table within exact provider season ID; special competition tie-breaks/formats not verified.",
        "history_scope": "Team form, H2H and table evidence are restricted to the exact provider season ID because a verified parent competition ID is unavailable; no cross-season or fuzzy competition joins are attempted.",
        "risk_policy": "Evidence coverage gap score only; not outcome risk or probability of losing.",
        "minimum_expected_league_groups": args.expected_league_groups,
        "league_group_gate": "PASS" if summary["selected_league_groups"] == args.expected_league_groups and report_ids.issubset(governed_ids) else "FAIL",
        "column_nonempty_counts": {h: sum(1 for r in report if r.get(h) not in ("", None)) for h in HEADERS},
        "required_join_gate": "PASS" if summary["unmatched_or_quarantined_rows"] == 0 and summary["duplicate_archive_composite_keys"] == 0 else "BLOCKED",
        "all_rows_research_only": True,
        "report_status": "PASS_RESEARCH_ONLY" if summary["unmatched_or_quarantined_rows"] == 0 and summary["selected_league_groups"] == args.expected_league_groups and report_ids.issubset(governed_ids) else "PARTIAL_RESEARCH_ONLY",
    })
    write_csv(out_dir / "oracle_selected50_market_intelligence.csv", report, HEADERS)
    write_csv(out_dir / "oracle_selected50_market_intelligence_evidence_trace.csv", trace,
              list(dict.fromkeys(k for r in trace for k in r)))
    write_xlsx(out_dir / "oracle_selected50_market_intelligence.xlsx", report)
    (out_dir / "oracle_selected50_market_intelligence_qa.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(json.dumps({k: summary[k] for k in (
        "runtime_rows", "report_rows", "exact_archive_joins", "unmatched_or_quarantined_rows",
        "selected_league_groups", "duplicate_archive_composite_keys", "schema_status",
        "league_group_gate", "required_join_gate", "production_eligible", "betting_eligible"
    )}, indent=2))
    if summary["schema_status"] != "PASS_38_COLUMNS":
        raise SystemExit("FAIL: report schema contract")
    if summary["league_group_gate"] != "PASS":
        raise SystemExit("FAIL: selected 50-league completeness gate")
    if summary["required_join_gate"] != "PASS":
        raise SystemExit("FAIL: exact archive join or duplicate-key gate")


if __name__ == "__main__":
    main()
