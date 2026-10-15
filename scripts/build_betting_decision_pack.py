#!/usr/bin/env python3
"""Build a reporting-only 22-market betting decision companion; never changes Oracle probabilities."""
from pathlib import Path
from datetime import date, timedelta
import json, os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT=Path(__file__).resolve().parents[1]
ORACLE=ROOT/"oracle"
PRED=ORACLE/"research/four_model/outputs/four_model_live_predictions.csv"
VI=ORACLE/"reports/visual_intelligence.csv"
SCHEDULE=ORACLE/"data/fixtures/oracle_upcoming_fixtures.csv"
MANIFEST=ORACLE/"research/four_model/outputs/four_model_live_prediction_manifest.json"
OUTDIR=ORACLE/"reports"
XLSX=OUTDIR/"oracle_betting_decision_pack.xlsx"
PDF=OUTDIR/"oracle_betting_decision_guide.pdf"
AUDIT=OUTDIR/"oracle_betting_decision_audit.json"

MARKETS=[
("BTTS_YES","BTTS (Yes)"),("BTTS_NO","BTTS (No)"),("HOME_1X","Home or Draw (1X)"),
("AWAY_X2","Away or Draw (X2)"),("HOME_12","No Draw (1 or 2)"),("OVER_1_5","Over 1.5 Goals"),
("OVER_2_5","Over 2.5 Goals"),("OVER_3_5","Over 3.5 Goals"),("UNDER_1_5","Under 1.5 Goals"),
("UNDER_2_5","Under 2.5 Goals"),("UNDER_3_5","Under 3.5 Goals"),("UNDER_4_5","Under 4.5 Goals"),
("HOME_O05","Home Team Over 0.5 Goals"),("HOME_O15","Home Team Over 1.5 Goals"),
("AWAY_O05","Away Team Over 0.5 Goals"),("AWAY_O15","Away Team Over 1.5 Goals"),
("HOME_1X_O15","Home 1X + Over 1.5"),("AWAY_X2_O15","Away X2 + Over 1.5"),
("HOME_1X_U45","Home 1X + Under 4.5"),("AWAY_X2_U45","Away X2 + Under 4.5"),
("M1UP_HOME","M-1UP Home"),("M1UP_AWAY","M-1UP Away")]

def report_window():
    start=os.getenv("REPORT_START_DATE","").strip()
    end=os.getenv("REPORT_END_DATE","").strip()
    if not start and not end:
        start=(date.today()+timedelta(days=1)).isoformat()
        end=(date.today()+timedelta(days=4)).isoformat()
    elif not start or not end:
        raise SystemExit("BLOCKED: provide both REPORT_START_DATE and REPORT_END_DATE in YYYY-MM-DD format.")
    try:
        sd=date.fromisoformat(start); ed=date.fromisoformat(end)
    except ValueError:
        raise SystemExit("BLOCKED: report dates must use YYYY-MM-DD.")
    if ed < sd: raise SystemExit("BLOCKED: report end date is earlier than start date.")
    return sd.isoformat(),ed.isoformat()

def main():
    start,end=report_window()
    for p in (PRED,VI,SCHEDULE,MANIFEST):
        if not p.exists() or p.stat().st_size==0: raise SystemExit(f"BLOCKED: required input missing/empty: {p}")
    pred=pd.read_csv(PRED).fillna("")
    vi=pd.read_csv(VI).fillna("")
    schedule=pd.read_csv(SCHEDULE).fillna("")
    manifest=json.loads(MANIFEST.read_text())
    required={"fixture_id","fixture_date","kickoff_utc","league","home_team","away_team","selection","raw_probability","status","calibration_status","data_cutoff_date"}
    if not required.issubset(pred.columns): raise SystemExit(f"BLOCKED: prediction schema missing {sorted(required-set(pred.columns))}")
    if "Fixture ID" not in vi.columns: raise SystemExit("BLOCKED: Visual Intelligence missing Fixture ID")
    if vi["Fixture ID"].isna().any() or vi["Fixture ID"].duplicated().any(): raise SystemExit("BLOCKED: Visual Intelligence Fixture IDs blank or duplicated")
    dates=pd.to_datetime(pred["fixture_date"],errors="coerce",utc=True).dt.strftime("%Y-%m-%d")
    pred=pred.loc[dates.between(start,end)].copy()
    if pred.empty: raise SystemExit(f"BLOCKED: no eligible predictions for {start} through {end}. Check ingestion window and dates.")
    if pred["fixture_id"].isna().any() or pred["fixture_id"].astype(str).str.strip().eq("").any(): raise SystemExit("BLOCKED: blank Oracle Fixture ID")
    identity_cols=["fixture_id","fixture_date","kickoff_utc","league","home_team","away_team","data_cutoff_date","state_date_home","state_date_away","training_n","source","status","calibration_status"]
    identities=pred[identity_cols].drop_duplicates()
    if identities.fixture_id.duplicated().any(): raise SystemExit("BLOCKED: fixture identity differs across market rows")
    present=pred.groupby("fixture_id")["selection"].agg(lambda x:set(x.astype(str)))
    expected={m for m,_ in MARKETS}
    bad={k:sorted(expected-v) for k,v in present.items() if v!=expected}
    if bad: raise SystemExit(f"BLOCKED: incomplete 22-market rows for fixtures: {list(bad.items())[:10]}")
    report_ids=set(identities.fixture_id.astype(str))
    vi=vi[vi["Fixture ID"].astype(str).isin(report_ids)].copy()
    if set(vi["Fixture ID"].astype(str)) != report_ids:
        raise SystemExit("BLOCKED: Visual Intelligence exact Fixture ID join is incomplete.")
    if len(vi)!=len(report_ids): raise SystemExit("BLOCKED: duplicate or excess Visual Intelligence rows after filtering.")
    pivot=pred.pivot(index="fixture_id",columns="selection",values="raw_probability").reset_index()
    fixtures=identities.merge(pivot,on="fixture_id",how="left",validate="one_to_one")
    fixtures=fixtures.merge(vi,left_on="fixture_id",right_on="Fixture ID",how="left",validate="one_to_one")
    if len(fixtures)!=len(identities) or fixtures.fixture_id.duplicated().any(): raise SystemExit("BLOCKED: fixture count changed or duplicate fixture after exact join.")
    fixtures["Kickoff Evidence"]=fixtures["kickoff_utc"].astype(str).str.strip().map(lambda x:"VERIFIED UTC" if x not in ("","nan","None","NaT") else "UNVERIFIED / MISSING")
    fixtures["Prematch Cutoff Check"]=[
        "PASS" if str(c) not in ("","nan","None") and str(d) not in ("","nan","None") and pd.to_datetime(c,errors="coerce") < pd.to_datetime(d,errors="coerce") else "CHECK REQUIRED"
        for c,d in zip(fixtures["data_cutoff_date"],fixtures["fixture_date"])]
    fixtures["Raw Probability Status"]="RAW / NOT CALIBRATED"
    fixtures["Bookmaker Odds Available"]="NO — not included in source artifact"
    fixtures["Value Assessment"]="NOT ASSESSED — bookmaker odds unavailable"
    fixtures["Staking Recommendation"]="NO AUTOMATIC STAKE — value and validation gates incomplete"
    fixtures["Betting Decision Gate"]=[
        "REVIEWABLE — ODDS/VALUE STILL REQUIRED"
        if ko=="VERIFIED UTC" and es=="VERIFIED_STANDINGS_AND_HISTORY" and pd.to_numeric(h,errors="coerce")>=3 and pd.to_numeric(a,errors="coerce")>=3
        else "CAUTION / INCOMPLETE EVIDENCE"
        for ko,es,h,a in zip(fixtures["Kickoff Evidence"],fixtures.get("Evidence Status",pd.Series([""]*len(fixtures))),fixtures.get("Home History Matches",pd.Series([0]*len(fixtures))),fixtures.get("Away History Matches",pd.Series([0]*len(fixtures)))]
    market_labels=dict(MARKETS)
    detail=pred[identity_cols+["selection","raw_probability","model_family","target_definition"]].copy()
    detail["Market"]=detail["selection"].map(market_labels)
    detail["Model Probability % (raw)"]=pd.to_numeric(detail["raw_probability"],errors="coerce")*100
    detail["Model-Implied Decimal Odds (1/p; raw only)"]=1/pd.to_numeric(detail["raw_probability"],errors="coerce").replace(0,float("nan"))
    detail["Bookmaker Odds (enter manually)"]=""
    detail["Raw EV Diagnostic (p*odds - 1)"]=""
    detail["Required Odds for 5% Raw Edge"]=1.05/pd.to_numeric(detail["raw_probability"],errors="coerce").replace(0,float("nan"))
    lookup=fixtures.set_index("fixture_id")
    for out_col,source_col in [("Evidence Status","Evidence Status"),("Evidence Risk","Risk Profile"),("Home History Matches","Home History Matches"),("Away History Matches","Away History Matches")]:
        detail[out_col]=detail["fixture_id"].map(lookup[source_col]) if source_col in lookup.columns else "UNAVAILABLE"
    detail["Kickoff Evidence"]=detail["fixture_id"].map(lookup["Kickoff Evidence"])
    detail["Pre-match Cutoff Check"]=detail["fixture_id"].map(lookup["Prematch Cutoff Check"])
    detail["First Goal / 1UP Validation"]=detail["selection"].map(lambda x:"NOT VERIFIED BY FINAL SCORES; event timestamps required" if str(x).startswith("M1UP") else "Not applicable")
    detail["Value Gate"]="BLOCKED: no bookmaker odds + probabilities uncalibrated"
    detail["Decision"]="NO BET RECOMMENDATION: require verified odds, price edge, and evidence checks"
    detail=detail.sort_values(["fixture_date","league","home_team","selection"])
    ex=pd.DataFrame(manifest.get("exclusion_details",[]))
    if len(ex):
        ex["fixture_date"]=ex["fixture_id"].astype(str).str.extract(r"SCHED-[^-]+-(\d{4}-\d{2}-\d{2})",expand=False)
        ex=ex[ex["fixture_date"].astype(str).between(start,end)].copy()
        ex["Next Action"]="Repair/verify team alias and role-history data; rerun before considering a prediction"
        ex_cols=[c for c in ["fixture_id","fixture_date","league","home_team","away_team","canonical_home_team","canonical_away_team","reason","missing_components","home_role_state_rows","away_role_state_rows","latest_home_role_state_date","latest_away_role_state_date","Next Action"] if c in ex.columns]
        ex=ex[ex_cols]
    else: ex=pd.DataFrame(columns=["fixture_id","fixture_date","league","home_team","away_team","reason","missing_components","Next Action"])
    scheduled=schedule[schedule["date"].astype(str).between(start,end)]
    evidence_counts=fixtures.get("Evidence Status",pd.Series(["UNAVAILABLE"]*len(fixtures))).fillna("UNAVAILABLE").value_counts().to_dict()
    kickoff_verified=int((fixtures["Kickoff Evidence"]=="VERIFIED UTC").sum())
    cutoff_pass=int((fixtures["Prematch Cutoff Check"]=="PASS").sum())
    audit_rows=[
        ("Report window",f"{start} through {end}"),("Configured universe","27 governed leagues; only leagues with fixtures in this window are shown"),
        ("Eligible predicted fixtures",int(fixtures.fixture_id.nunique())),("Markets per eligible fixture",22),("Fixture-market rows",int(len(detail))),
        ("Scheduled fixtures in window",int(len(scheduled))),("Excluded fixture records in window",int(len(ex))),
        ("Kickoff verified UTC",kickoff_verified),("Kickoff missing/unverified",int(len(fixtures)-kickoff_verified)),
        ("Evidence status counts",json.dumps(evidence_counts)),("Pre-match cutoff check passed",f"{cutoff_pass}/{len(fixtures)}"),
        ("Exact Fixture ID join","PASS"),("Oracle probabilities modified",False),("Calibration","NOT_APPLIED_RAW_ONLY"),
        ("Production promotion","BLOCKED"),("xG included",False),("Bookmaker odds","Unavailable in source artifact; no EV conclusion"),
        ("M-1UP validation","Requires event-level goal timestamps; final scores insufficient"),
        ("Historical score audit","Research-only challenger audit; not production accuracy; canonical settlement ledger empty"),
        ("Interpretation","Raw probability and evidence coverage are not a guarantee of outcome or value.")
    ]
    coverage=pd.DataFrame(audit_rows,columns=["Audit Item","Result"])
    guide=pd.DataFrame([
        ("Probability","All Oracle probabilities are raw and uncalibrated; do not treat as verified event frequencies."),
        ("Fair odds","1/p is a mathematical raw-probability equivalent only, not a validated price target."),
        ("Bookmaker odds","Enter current decimal odds in Market Detail to enable a raw EV diagnostic; price/value still requires calibrated probabilities."),
        ("Raw EV","p × odds − 1 is only a screening diagnostic while calibration is NOT_APPLIED."),
        ("5% raw edge hurdle","Required odds = 1.05/p; screening threshold only, not a bet recommendation."),
        ("Evidence risk","Risk Profile is evidence coverage, not outcome safety. Missing evidence is not a negative team signal."),
        ("M-1UP","First-goal/lead truth needs event-level timestamps; final scores cannot establish 1UP success."),
        ("Governance","Companion pack only. Locked 32-column report, Oracle production logic, raw probabilities and promotion gates remain unchanged; xG excluded."),
        ("No automatic stakes","No selections are labeled confirmed value bets; bookmaker odds and validated market-specific performance are required.")
    ],columns=["Topic","Rule / Interpretation"])
    OUTDIR.mkdir(parents=True,exist_ok=True)
    xlsx_tmp=XLSX
    with pd.ExcelWriter(xlsx_tmp,engine="openpyxl") as writer:
        guide.to_excel(writer,index=False,sheet_name="Read Me - Decision Gates")
        fixtures.to_excel(writer,index=False,sheet_name="Fixture Overview - 22 Markets")
        detail.to_excel(writer,index=False,sheet_name="Market Detail - Odds Value")
        ex.to_excel(writer,index=False,sheet_name="Excluded Fixtures")
        coverage.to_excel(writer,index=False,sheet_name="Coverage Audit")
    wb=load_workbook(XLSX)
    header_fill=PatternFill("solid",fgColor="1F4E78"); header_font=Font(color="FFFFFF",bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws.sheet_view.showGridLines=False
        for c in ws[1]:
            c.fill=header_fill;c.font=header_font;c.alignment=Alignment(vertical="top",wrap_text=True)
        ws.row_dimensions[1].height=32
        for ci,cells in enumerate(ws.columns,1):
            vals=[str(c.value) if c.value is not None else "" for c in list(cells)[:120]]
            width=min(max([len(v) for v in vals]+[10]),42)
            ws.column_dimensions[get_column_letter(ci)].width=max(12,width+2)
        for row in ws.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(vertical="top",wrap_text=False)
    ws=wb["Fixture Overview - 22 Markets"]
    hdr=[c.value for c in ws[1]]
    for col in [m for m,_ in MARKETS]:
        if col in hdr:
            idx=hdr.index(col)+1
            for rr in range(2,ws.max_row+1):ws.cell(rr,idx).number_format="0.0%"
    ws=wb["Market Detail - Odds Value"];hdr=[c.value for c in ws[1]]
    for col in ["raw_probability"]:
        idx=hdr.index(col)+1
        for rr in range(2,ws.max_row+1):ws.cell(rr,idx).number_format="0.0%"
    for col in ["Model-Implied Decimal Odds (1/p; raw only)","Required Odds for 5% Raw Edge"]:
        idx=hdr.index(col)+1
        for rr in range(2,ws.max_row+1):ws.cell(rr,idx).number_format="0.00"
    p_idx=hdr.index("raw_probability")+1
    odds_idx=hdr.index("Bookmaker Odds (enter manually)")+1
    ev_idx=hdr.index("Raw EV Diagnostic (p*odds - 1)")+1
    for rr in range(2,ws.max_row+1):
        pcell=f"{get_column_letter(p_idx)}{rr}"
        ocell=f"{get_column_letter(odds_idx)}{rr}"
        ws.cell(rr,ev_idx).value=f'=IF(OR({ocell}="",{ocell}<=1),"",{pcell}*{ocell}-1)'
        ws.cell(rr,ev_idx).number_format="0.0%"
    wb.save(XLSX)
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SmallBody",parent=styles["BodyText"],fontSize=8,leading=10,spaceAfter=4))
    styles.add(ParagraphStyle(name="SmallHeading",parent=styles["Heading2"],fontSize=11,leading=13,spaceBefore=7,spaceAfter=4))
    pdf=SimpleDocTemplate(str(PDF),pagesize=landscape(A4),rightMargin=28,leftMargin=28,topMargin=26,bottomMargin=26)
    story=[Paragraph("Oracle Betting Decision Guide",styles["Title"]),
      Paragraph(f"Companion to the locked Oracle Curated Match Report | Fixtures: {start} to {end}",styles["SmallBody"]),
      Paragraph("The companion workbook contains all 22 governed market probabilities per eligible fixture, Visual Intelligence evidence, excluded fixtures, and a market-level worksheet for recording bookmaker odds.",styles["SmallBody"]),
      Paragraph("Decision-critical status",styles["SmallHeading"])]
    table_data=[["Item","Current status"],
      ["Eligible fixtures",f"{len(fixtures)} fixtures; {len(detail)} fixture-market rows"],
      ["Scheduled fixtures in window",str(len(scheduled))],
      ["Excluded fixtures",str(len(ex))+" — missing pre-match team-state records; details in workbook"],
      ["Kickoff verified UTC",f"{kickoff_verified}/{len(fixtures)}"],
      ["Kickoff missing/unverified",str(len(fixtures)-kickoff_verified)],
      ["Evidence status",json.dumps(evidence_counts)],
      ["Odds/value","Bookmaker odds not present; value cannot be assessed"],
      ["Calibration / promotion","Raw only; calibration NOT_APPLIED; production promotion BLOCKED"],
      ["M-1UP truth","Requires event-level goal timestamps"]]
    t=Table(table_data,colWidths=[185,530],repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1F4E78")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
      ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),8),("VALIGN",(0,0),(-1,-1),"TOP"),
      ("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#D9E2F3")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F3F6FA")]),
      ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
    story.extend([t,Spacer(1,8),Paragraph("Use the Market Detail sheet to enter actual bookmaker decimal odds. The raw EV field is a screening diagnostic only while Oracle probabilities remain uncalibrated. High probability is not automatically value; this report does not provide automatic stakes or confirm any value bet.",styles["SmallBody"]),
      Paragraph("Governance: the locked 32-column report and production model are not modified; raw probabilities are preserved, calibration remains NOT_APPLIED, promotion remains BLOCKED, and xG remains excluded.",styles["SmallBody"])])
    pdf.build(story)
    audit={"status":"PASS","report_window":{"start":start,"end":end},"fixture_count":int(len(fixtures)),"markets_per_fixture":22,
      "fixture_market_rows":int(len(detail)),"scheduled_fixture_count":int(len(scheduled)),"excluded_fixture_count":int(len(ex)),
      "visual_intelligence_exact_join":True,"duplicate_fixture_ids":0,"oracle_probabilities_modified":False,
      "calibration":"NOT_APPLIED","production_promotion":"BLOCKED","xg_included":False,
      "kickoff_coverage":{"verified_utc":kickoff_verified,"missing_or_unverified":int(len(fixtures)-kickoff_verified)},
      "evidence_status_counts":evidence_counts,"prematch_cutoff_check_passed":cutoff_pass,
      "bookmaker_odds_available":False,"value_assessed":False,"m1up_event_truth_verified":False,
      "workbook":str(XLSX),"pdf":str(PDF)}
    AUDIT.write_text(json.dumps(audit,indent=2,default=str)+"\n")
    print(json.dumps(audit,indent=2,default=str))

if __name__=="__main__":main()
