import importlib.util
import unittest
import tempfile
from pathlib import Path
from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "build_selected50_market_intelligence_report.py"
spec = importlib.util.spec_from_file_location("market_report", SCRIPT)
market_report = importlib.util.module_from_spec(spec)
spec.loader.exec_module(market_report)


def archive_row(fid, date, home_id, away_id, hg, ag, season="S1", league="L1"):
    return {
        "provider_season_id": season, "fixture_id": fid,
        "home_team_id": home_id, "away_team_id": away_id,
        "kickoff_utc": date, "league_folder": league, "season_folder": "2026",
        "status_raw": "complete", "home_goals": str(hg), "away_goals": str(ag),
        "home_team": home_id, "away_team": away_id,
    }


def runtime_row():
    return {
        "fixture_id": "F100", "date_utc": "2026-10-31T18:00:00+00:00",
        "selected_id": "S1", "provider_season_id": "S1",
        "home_team_provider_id": "H", "away_team_provider_id": "A",
        "league": "Test League", "home_team": "Home FC", "away_team": "Away FC",
        "BTTS_probability_pct": "58.21", "DC_HOME_probability_pct": "71.09",
        "DC_AWAY_probability_pct": "49.87",
    }


class Selected50MarketIntelligenceTests(unittest.TestCase):
    def setUp(self):
        self.archive = [
            archive_row("F1", "2026-10-01T12:00:00+00:00", "H", "X", 2, 1),
            archive_row("F2", "2026-10-03T12:00:00+00:00", "Y", "A", 0, 1),
            archive_row("F3", "2026-10-08T12:00:00+00:00", "H", "Z", 1, 1),
            archive_row("F4", "2026-10-10T12:00:00+00:00", "W", "A", 2, 0),
            archive_row("F5", "2026-10-15T12:00:00+00:00", "H", "A", 3, 1),
            archive_row("F6", "2026-10-20T12:00:00+00:00", "A", "Q", 1, 2),
            archive_row("F7", "2026-10-25T12:00:00+00:00", "H", "P", 2, 0),
            archive_row("F100", "2026-10-31T18:00:00+00:00", "H", "A", 0, 0),
        ]

    def test_exact_join_requires_season_fixture_and_ordered_team_ids(self):
        rows, missing, duplicates = market_report.exact_join([runtime_row()], self.archive)
        self.assertEqual(len(rows), 1)
        self.assertEqual(missing, [])
        swapped = dict(runtime_row(), home_team_provider_id="A", away_team_provider_id="H")
        rows, missing, _ = market_report.exact_join([swapped], self.archive)
        self.assertEqual(rows, {})
        self.assertEqual(missing[0][1], "NO_EXACT_ARCHIVE_MATCH")

    def test_kickoff_mismatch_is_quarantined(self):
        changed = dict(self.archive[-1], kickoff_utc="2026-10-31T17:59:00+00:00")
        joined, missing, _ = market_report.exact_join([runtime_row()], self.archive[:-1] + [changed])
        self.assertEqual(joined, {})
        self.assertEqual(missing[0][1], "KICKOFF_MISMATCH_OR_UNPARSEABLE")

    def test_same_kickoff_fixture_is_not_used_as_history(self):
        same_time = archive_row("F101", "2026-10-31T18:00:00+00:00", "H", "B", 4, 0)
        report, trace, summary = market_report.build_report([runtime_row()], self.archive + [same_time])
        self.assertEqual(trace[0]["home_form_n"], 4)
        self.assertEqual(trace[0]["away_form_n"], 4)
        self.assertEqual(summary["exact_archive_joins"], 1)

    def test_raw_probabilities_are_copied_and_unsupported_markets_masked(self):
        report, _, _ = market_report.build_report([runtime_row()], self.archive)
        row = report[0]
        self.assertEqual(row["BTTS Yes"], "58.21")
        self.assertEqual(row["Home 1X"], "71.09")
        self.assertEqual(row["Away X2"], "49.87")
        for header in market_report.HEADERS[6:20]:
            self.assertEqual(row[header], market_report.UNAVAILABLE)
        self.assertEqual(len(market_report.HEADERS), 38)
        self.assertFalse(any("COMBO" in h and row[h] not in ("", market_report.UNAVAILABLE)
                             for h in row))

    def test_xlsx_has_exact_locked_38_column_schema_and_research_only_trace(self):
        report, trace, _ = market_report.build_report([runtime_row()], self.archive)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            market_report.write_xlsx(path, report)
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook["Market Intelligence"]
            self.assertEqual([cell.value for cell in sheet[1]], market_report.HEADERS)
            self.assertEqual(sheet.max_column, 38)
            workbook.close()
        self.assertEqual(trace[0]["production_eligible"], "NO")
        self.assertEqual(trace[0]["betting_eligible"], "NO")
        self.assertEqual(trace[0]["history_cutoff_rule"], "STRICTLY_BEFORE_FIXTURE_KICKOFF")

    def test_duplicate_runtime_keys_fail_closed(self):
        with self.assertRaisesRegex(ValueError, "Duplicate runtime exact key"):
            market_report.build_report([runtime_row(), runtime_row()], self.archive)

    def test_duplicate_archive_key_is_not_silently_joined(self):
        duplicate = dict(self.archive[-1])
        joined, missing, duplicate_keys = market_report.exact_join(
            [runtime_row()], self.archive + [duplicate]
        )
        self.assertEqual(joined, {})
        self.assertEqual(missing[0][1], "DUPLICATE_ARCHIVE_MATCH")
        self.assertTrue(duplicate_keys)

    def test_missing_identity_is_quarantined_not_name_matched(self):
        row = dict(runtime_row(), provider_season_id="")
        joined, missing, _ = market_report.exact_join([row], self.archive)
        self.assertEqual(joined, {})
        self.assertEqual(missing[0][1], "MISSING_COMPOSITE_IDENTITY")


if __name__ == "__main__":
    unittest.main()
